#!/usr/bin/env python3
"""
Controle-export: producttypes uit Excel-tabblad Products, plus alle Shopify-producten
met die types, huidige prijs en prijs +9%.

Geen writes naar Shopify.

  python3 scripts/shopify_wp_type_markup_preview.py
  python3 scripts/shopify_wp_type_markup_preview.py --xlsx "/pad/naar/Aanpassing prijzen WP.xlsx"

Output (output/):
  wp_type_markup_types_<stamp>.csv
  wp_type_markup_producten_<stamp>.csv
"""

from __future__ import annotations

import argparse
import csv
import os
import sys
import time
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:
    print("Installeer openpyxl: pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)

try:
    import requests
except ImportError:
    print("Installeer requests: pip install requests", file=sys.stderr)
    raise SystemExit(1)

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

from modules.env_loader import load_project_env  # noqa: E402
from modules.price_markup import EXCLUDED_TYPES  # noqa: E402

load_project_env()

import config  # noqa: E402

DEFAULT_XLSX = Path("/Users/donolsthoorn/Downloads/Aanpassing prijzen WP.xlsx")
MARKUP = Decimal("1.09")
CENT = Decimal("0.01")
_REQUEST_TIMEOUT = (15, 120)

QUERY_PRODUCTS = """
query ($q: String, $c: String) {
  products(first: 50, after: $c, query: $q) {
    pageInfo { hasNextPage endCursor }
    nodes {
      id
      handle
      title
      vendor
      productType
      status
      tags
      variants(first: 100) {
        pageInfo { hasNextPage endCursor }
        nodes {
          id
          sku
          title
          barcode
          price
          compareAtPrice
        }
      }
    }
  }
}
"""

QUERY_MORE_VARIANTS = """
query ($id: ID!, $c: String) {
  product(id: $id) {
    variants(first: 100, after: $c) {
      pageInfo { hasNextPage endCursor }
      nodes {
        id
        sku
        title
        barcode
        price
        compareAtPrice
      }
    }
  }
}
"""

PRODUCT_FIELDS = [
    "Handle",
    "Product_ID",
    "Titel",
    "Vendor",
    "Type",
    "Status",
    "Tags",
    "Variant_ID",
    "SKU",
    "Variant_titel",
    "Barcode",
    "Huidige_prijs",
    "Compare_at_prijs",
    "Prijs_plus_9pct",
    "Prijs_plus_9pct_afgerond",
    "Verschil_afgerond",
    "Staat_in_excel",
]

TYPE_FIELDS = [
    "Type",
    "In_excel_rijen",
    "In_excel_unieke_handles",
    "Shop_producten",
    "Shop_varianten",
    "Vendors",
    "Statussen",
    "Voorbeeld_titel",
]


def _session() -> requests.Session:
    s = requests.Session()
    s.trust_env = False
    return s


def _gql(
    sess: requests.Session,
    shop: str,
    token: str,
    api: str,
    query: str,
    variables: dict | None = None,
) -> dict:
    url = f"https://{shop}/admin/api/{api}/graphql.json"
    last: dict = {}
    for attempt in range(12):
        r = sess.post(
            url,
            headers={"X-Shopify-Access-Token": token, "Content-Type": "application/json"},
            json={"query": query, "variables": variables or {}},
            timeout=_REQUEST_TIMEOUT,
            proxies={"http": None, "https": None},
        )
        if r.status_code == 429:
            time.sleep(1.5 + attempt)
            continue
        if r.status_code >= 500:
            time.sleep(3)
            continue
        r.raise_for_status()
        last = r.json()
        errs = last.get("errors") or []
        throttled = any((e.get("extensions") or {}).get("code") == "THROTTLED" for e in errs)
        if throttled:
            time.sleep(min(2.0 * (attempt + 1), 20.0))
            continue
        if errs:
            raise RuntimeError(errs)
        return last
    raise RuntimeError(last.get("errors") or "GraphQL mislukt")


def _gid_num(gid: str) -> str:
    return (gid or "").rsplit("/", 1)[-1]


def _norm_handle(raw: Any) -> str:
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return str(raw).lower()
    if isinstance(raw, int):
        return str(raw).lower()
    if isinstance(raw, float):
        if raw.is_integer():
            return str(int(raw)).lower()
        return str(raw).strip().lower()
    s = str(raw).strip().lower()
    if s.endswith(".0") and s[:-2].replace("-", "").isdigit():
        s = s[:-2]
    return s


def _handle_keys(handle: str) -> set[str]:
    h = _norm_handle(handle)
    if not h:
        return set()
    out = {h}
    for prefix in ("wp-", "hsq-"):
        if h.startswith(prefix) and len(h) > len(prefix):
            out.add(h[len(prefix) :])
    return out


def _in_excel(handle: str, excel_handles: set[str]) -> bool:
    return any(k in excel_handles for k in _handle_keys(handle))


def _dec(raw: Any) -> Decimal | None:
    if raw is None:
        return None
    s = str(raw).strip().replace(",", ".")
    if not s:
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _fmt(d: Decimal | None, places: int = 4) -> str:
    if d is None:
        return ""
    q = Decimal("0." + "0" * places) if places else Decimal("1")
    return format(d.quantize(q), f"f")


def _round_cent(d: Decimal) -> Decimal:
    return d.quantize(CENT, rounding=ROUND_HALF_UP)


def _search_query(product_type: str) -> str:
    escaped = product_type.replace("\\", "\\\\").replace('"', '\\"')
    return f'product_type:"{escaped}"'


def load_excel_types(path: Path) -> tuple[list[str], set[str], dict[str, int], dict[str, int]]:
    if not path.is_file():
        raise SystemExit(f"Excel niet gevonden: {path}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    name = wb.sheetnames[0]
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise SystemExit(f"Leeg tabblad: {name}")
    cols = {str(c or "").strip().lower(): i for i, c in enumerate(header)}
    type_i = cols.get("type")
    handle_i = cols.get("handle")
    if type_i is None:
        raise SystemExit(f"Kolom Type ontbreekt in tabblad {name}")
    type_order: list[str] = []
    seen_types: set[str] = set()
    type_rows: dict[str, int] = defaultdict(int)
    type_handles: dict[str, set[str]] = defaultdict(set)
    excel_handles: set[str] = set()
    for row in rows:
        if not row or type_i >= len(row):
            continue
        ptype = str(row[type_i] or "").strip()
        if not ptype:
            continue
        if ptype not in seen_types:
            seen_types.add(ptype)
            type_order.append(ptype)
        type_rows[ptype] += 1
        if handle_i is not None and handle_i < len(row):
            h = _norm_handle(row[handle_i])
            if h:
                excel_handles.add(h)
                type_handles[ptype].add(h)
    wb.close()
    handle_counts = {t: len(type_handles[t]) for t in type_order}
    print(f"Excel tabblad '{name}': {len(type_order)} types, {sum(type_rows.values())} rijen", flush=True)
    return type_order, excel_handles, dict(type_rows), handle_counts


def _all_variants(
    sess: requests.Session,
    shop: str,
    token: str,
    api: str,
    product: dict[str, Any],
) -> list[dict[str, Any]]:
    conn = product.get("variants") or {}
    nodes = list(conn.get("nodes") or [])
    info = conn.get("pageInfo") or {}
    cursor = info.get("endCursor")
    while info.get("hasNextPage"):
        body = _gql(
            sess,
            shop,
            token,
            api,
            QUERY_MORE_VARIANTS,
            {"id": product.get("id"), "c": cursor},
        )
        conn = ((body.get("data") or {}).get("product") or {}).get("variants") or {}
        nodes.extend(conn.get("nodes") or [])
        info = conn.get("pageInfo") or {}
        cursor = info.get("endCursor")
        time.sleep(0.2)
    return nodes


def fetch_type_products(
    sess: requests.Session,
    shop: str,
    token: str,
    api: str,
    product_type: str,
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    cursor = None
    q = _search_query(product_type)
    while True:
        body = _gql(sess, shop, token, api, QUERY_PRODUCTS, {"q": q, "c": cursor})
        conn = ((body.get("data") or {}).get("products") or {})
        for p in conn.get("nodes") or []:
            shop_type = str(p.get("productType") or "").strip()
            if shop_type.casefold() != product_type.casefold():
                continue
            variants = _all_variants(sess, shop, token, api, p)
            out.append({**p, "_variants": variants})
        info = conn.get("pageInfo") or {}
        if not info.get("hasNextPage"):
            break
        cursor = info.get("endCursor")
        time.sleep(0.25)
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="Preview +9% prijs per type uit Excel")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX, help="Excel met types in 1e tabblad")
    ap.add_argument(
        "--out-dir",
        type=Path,
        default=ROOT / "output",
        help="Outputmap",
    )
    args = ap.parse_args()

    shop = (config.SHOPIFY_SHOP_DOMAIN or "").strip()
    token = (config.SHOPIFY_ACCESS_TOKEN or "").strip()
    api = (config.SHOPIFY_ADMIN_API_VERSION or "2024-10").strip()
    if not shop or not token:
        print("SHOPIFY_SHOP_DOMAIN / SHOPIFY_ACCESS_TOKEN ontbreekt", file=sys.stderr)
        return 2

    types, excel_handles, excel_row_counts, excel_handle_counts = load_excel_types(args.xlsx)
    skipped = [t for t in types if t.casefold() in {x.casefold() for x in EXCLUDED_TYPES}]
    types = [t for t in types if t.casefold() not in {x.casefold() for x in EXCLUDED_TYPES}]
    if skipped:
        print("Uitgesloten types (geen +9%):", flush=True)
        for t in skipped:
            print(f"  - {t}", flush=True)
    sess = _session()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    args.out_dir.mkdir(parents=True, exist_ok=True)
    products_path = args.out_dir / f"wp_type_markup_producten_{stamp}.csv"
    types_path = args.out_dir / f"wp_type_markup_types_{stamp}.csv"

    product_rows: list[dict[str, str]] = []
    type_rows: list[dict[str, str]] = []
    seen_products: set[str] = set()

    print(f"Shop: {shop}", flush=True)
    for i, ptype in enumerate(types, start=1):
        print(f"[{i}/{len(types)}] {ptype}…", flush=True)
        products = fetch_type_products(sess, shop, token, api, ptype)
        vendors: dict[str, int] = defaultdict(int)
        statuses: dict[str, int] = defaultdict(int)
        n_var = 0
        sample_title = ""
        for p in products:
            pid = _gid_num(p.get("id") or "")
            if pid in seen_products:
                continue
            seen_products.add(pid)
            handle = str(p.get("handle") or "").strip()
            title = str(p.get("title") or "").strip()
            vendor = str(p.get("vendor") or "").strip()
            status = str(p.get("status") or "").strip()
            tags = ", ".join(p.get("tags") or [])
            vendors[vendor or "(leeg)"] += 1
            statuses[status or "(leeg)"] += 1
            if not sample_title:
                sample_title = title
            in_excel = "ja" if _in_excel(handle, excel_handles) else "nee"
            variants = p.get("_variants") or []
            if not variants:
                n_var += 1
                product_rows.append(
                    {
                        "Handle": handle,
                        "Product_ID": pid,
                        "Titel": title,
                        "Vendor": vendor,
                        "Type": ptype,
                        "Status": status,
                        "Tags": tags,
                        "Variant_ID": "",
                        "SKU": "",
                        "Variant_titel": "",
                        "Barcode": "",
                        "Huidige_prijs": "",
                        "Compare_at_prijs": "",
                        "Prijs_plus_9pct": "",
                        "Prijs_plus_9pct_afgerond": "",
                        "Verschil_afgerond": "",
                        "Staat_in_excel": in_excel,
                    }
                )
                continue
            for v in variants:
                n_var += 1
                price = _dec(v.get("price"))
                plus = (price * MARKUP) if price is not None else None
                plus_r = _round_cent(plus) if plus is not None else None
                diff = (plus_r - price) if plus_r is not None and price is not None else None
                product_rows.append(
                    {
                        "Handle": handle,
                        "Product_ID": pid,
                        "Titel": title,
                        "Vendor": vendor,
                        "Type": ptype,
                        "Status": status,
                        "Tags": tags,
                        "Variant_ID": _gid_num(v.get("id") or ""),
                        "SKU": str(v.get("sku") or "").strip(),
                        "Variant_titel": str(v.get("title") or "").strip(),
                        "Barcode": str(v.get("barcode") or "").strip(),
                        "Huidige_prijs": _fmt(price, 2) if price is not None else "",
                        "Compare_at_prijs": _fmt(_dec(v.get("compareAtPrice")), 2),
                        "Prijs_plus_9pct": _fmt(plus, 4),
                        "Prijs_plus_9pct_afgerond": _fmt(plus_r, 2),
                        "Verschil_afgerond": _fmt(diff, 2),
                        "Staat_in_excel": in_excel,
                    }
                )
        type_rows.append(
            {
                "Type": ptype,
                "In_excel_rijen": str(excel_row_counts.get(ptype, 0)),
                "In_excel_unieke_handles": str(excel_handle_counts.get(ptype, 0)),
                "Shop_producten": str(len(products)),
                "Shop_varianten": str(n_var),
                "Vendors": ", ".join(f"{k}={v}" for k, v in sorted(vendors.items())),
                "Statussen": ", ".join(f"{k}={v}" for k, v in sorted(statuses.items())),
                "Voorbeeld_titel": sample_title,
            }
        )
        print(f"  → {len(products)} producten, {n_var} varianten", flush=True)
        time.sleep(0.2)

    with types_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=TYPE_FIELDS, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        w.writerows(type_rows)
    with products_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=PRODUCT_FIELDS, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        w.writerows(product_rows)

    n_prod = len({r["Product_ID"] for r in product_rows})
    n_new = sum(1 for r in product_rows if r["Staat_in_excel"] == "nee")
    vendors_all = sorted({r["Vendor"] for r in product_rows if r["Vendor"]})
    print(
        f"\nKlaar: {len(types)} types, {n_prod} producten, {len(product_rows)} variant-rijen.",
        flush=True,
    )
    print(f"Niet in Excel-tabblad (wel zelfde type): {n_new} rijen", flush=True)
    print(f"Vendors: {', '.join(vendors_all)}", flush=True)
    print(f"Types: {types_path}", flush=True)
    print(f"Producten: {products_path}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
